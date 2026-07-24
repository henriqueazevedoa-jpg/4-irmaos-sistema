#!/usr/bin/env python3
"""
Gera uma planilha (revisao-produtos.xlsx) só com os casos que precisam de
olho humano depois da faxina automática (Fase 2). Abas:
  - Duplicatas p/ revisar: mesmo produto com marca/escrita diferente
  - Sem preço: itens que ficaram sem preço de venda
  - Preços p/ conferir: itens caros (podem ser pacote ou erro de custo)

Uso: python3 gerar-revisao.py produtos-import.json saida.xlsx
"""
import sys, json, re
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font

CAB = Font(bold=True)


def base_fp(nome, marca):
    toks = re.sub(r"[^A-Z0-9 ]", " ", nome.upper()).split()
    marca_toks = set(re.sub(r"[^A-Z0-9 ]", " ", (marca or "").upper()).split())
    toks = [t for t in toks if t not in marca_toks]
    return "".join(sorted(toks))


def cabecalho(ws, colunas):
    ws.append(colunas)
    for c in ws[1]:
        c.font = CAB


def main():
    entrada = sys.argv[1] if len(sys.argv) > 1 else "produtos-import.json"
    saida = sys.argv[2] if len(sys.argv) > 2 else "revisao-produtos.xlsx"
    produtos = json.load(open(entrada, encoding="utf-8"))["produtos"]

    wb = Workbook()

    # 1) Duplicatas p/ revisar (mesmo produto-base, escrita/marca diferente)
    grupos = defaultdict(list)
    for p in produtos:
        grupos[base_fp(p["descricao"], p.get("fabricante"))].append(p)
    dups = {k: v for k, v in grupos.items() if len(v) > 1 and len({x["descricao"] for x in v}) > 1}

    ws = wb.active
    ws.title = "Duplicatas p revisar"
    cabecalho(ws, ["Grupo", "Nome", "Marca", "Custo", "Preço", "Estoque", "Código"])
    for i, (k, v) in enumerate(sorted(dups.items()), 1):
        for p in v:
            ws.append([i, p["descricao"], p.get("fabricante") or "", p["precoCusto"],
                       p["precoVenda"], p["saldoEstoque"], p.get("codigoExterno") or ""])

    # 2) Sem preço de venda
    ws2 = wb.create_sheet("Sem preço")
    cabecalho(ws2, ["Nome", "Marca", "Custo", "Estoque", "Código"])
    sem = [p for p in produtos if not p["precoVenda"]]
    for p in sem:
        ws2.append([p["descricao"], p.get("fabricante") or "", p["precoCusto"],
                    p["saldoEstoque"], p.get("codigoExterno") or ""])

    # 3) Preços p/ conferir (caros — pode ser pacote ou erro)
    ws3 = wb.create_sheet("Precos p conferir")
    cabecalho(ws3, ["Nome", "Marca", "Custo", "Preço", "Estoque", "Código"])
    caros = sorted([p for p in produtos if p["precoVenda"] > 200], key=lambda x: -x["precoVenda"])
    for p in caros:
        ws3.append([p["descricao"], p.get("fabricante") or "", p["precoCusto"],
                    p["precoVenda"], p["saldoEstoque"], p.get("codigoExterno") or ""])

    wb.save(saida)
    print(f"Planilha: {saida}")
    print(f"  Duplicatas p/ revisar: {len(dups)} grupos ({sum(len(v) for v in dups.values())} itens)")
    print(f"  Sem preço: {len(sem)}")
    print(f"  Preços p/ conferir (> R$200): {len(caros)}")


if __name__ == "__main__":
    main()
