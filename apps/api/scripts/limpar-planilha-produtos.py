#!/usr/bin/env python3
"""
Lê a planilha de produtos do sistema antigo (SIC) e gera um JSON normalizado
para importar. Faz a FAXINA (Fase 1):
  - limpa lixo do nome (início/fim), espaço duplo, padroniza MAIÚSCULO
  - reconstrói o preço corrompido pela exportação (custo x markup)
  - extrai a MARCA (coluna Fabricante + do próprio nome)
  - remove duplicatas óbvias (mesmo nome após normalizar), mantendo o de maior estoque
  - inicializa o custo médio com o último custo (coluna 'Custo médio' é corrompida)

Uso: python3 limpar-planilha-produtos.py "/caminho/planilha.xlsx" saida.json
"""
import sys, json, re
from collections import Counter
from openpyxl import load_workbook

UNI_MAP = {"UND": "UN", "UNID": "UN", "UNIDADE": "UN", "UN.": "UN", "PC.": "PC", "PÇ": "PC", "PCS": "PC"}


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s == "":
        return None
    if "," in s:  # formato brasileiro
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def code_str(v):
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, (int, float)):
        return str(int(v))
    return str(v).strip()


def ncm_clean(v):
    if v is None:
        return None
    s = "".join(ch for ch in str(v) if ch.isdigit())
    if not s:
        return None
    return s.zfill(8)[:8] if len(s) <= 8 else s[:8]


def barcode_clean(v):
    if v is None:
        return None
    s = str(v).strip().upper()
    if "SEM GTIN" in s or s == "":
        return None
    d = "".join(ch for ch in s if ch.isdigit())
    return d if 8 <= len(d) <= 14 else None


def unidade_clean(v):
    if v is None or str(v).strip() == "":
        return "UN"
    s = str(v).strip().upper()
    return UNI_MAP.get(s, s)[:10]


def limpar_nome(s):
    """Tira lixo do início/fim, colapsa espaços e padroniza maiúsculo.
    Preserva pontuação do meio (tamanhos como 1.1/2, 2,5X100, 50%)."""
    n = str(s).strip()
    n = re.sub(r"^[\s/.\-*>`~^#%]+", "", n)   # lixo no começo
    n = re.sub(r"[\s*>`~^#%.\-]+$", "", n)      # lixo no fim
    n = re.sub(r"\s+", " ", n)                   # espaço duplo
    return n.upper().strip()


def fingerprint(nome):
    """Chave para deduplicar: só letras/números (ignora espaço e pontuação)."""
    return re.sub(r"[^A-Z0-9]", "", nome.upper())


def extrair_marca(nome, fab_val, marcas):
    # 1) coluna Fabricante, se parecer uma marca (tem letras)
    if fab_val:
        f = str(fab_val).strip().upper()
        if len(f) >= 2 and any(c.isalpha() for c in f):
            return f
    # 2) do próprio nome: última palavra que seja marca conhecida
    toks = re.sub(r"[^A-Z0-9 ]", " ", nome.upper()).split()
    for t in reversed(toks):
        if t in marcas:
            return t
    return None


def texto(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def main():
    entrada = sys.argv[1]
    saida = sys.argv[2] if len(sys.argv) > 2 else "produtos-import.json"
    wb = load_workbook(entrada, read_only=True, data_only=True)
    ws = wb["Plan1"]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    idx = {h: i for i, h in enumerate(header)}
    data = [r for r in it if r[idx["Produto"]]]

    def g(r, name):
        return r[idx[name]] if name in idx else None

    # Conjunto de marcas conhecidas: fabricantes usados >= 5 vezes, com nome limpo
    fab = Counter(str(g(r, "Fabricante")).strip().upper() for r in data if g(r, "Fabricante"))
    marcas = {
        m for m, c in fab.items()
        if c >= 5 and 1 <= len(m.split()) <= 2 and len(m) >= 3 and m.replace(" ", "").isalpha()
    }

    produtos = []
    fornecedores = set()
    vistos_codigo = set()
    vistos_barras = set()
    pulados_dup = 0

    for r in data:
        descricao = limpar_nome(g(r, "Produto"))
        if not descricao:
            continue

        codigo = code_str(g(r, "Código"))
        if codigo and codigo in vistos_codigo:
            pulados_dup += 1
            continue
        if codigo:
            vistos_codigo.add(codigo)

        barras = barcode_clean(g(r, "Cean")) or barcode_clean(g(r, "Ceantrib"))
        if barras and barras in vistos_barras:
            barras = None
        if barras:
            vistos_barras.add(barras)

        fornecedor = texto(g(r, "Fornecedor"))
        if fornecedor:
            fornecedores.add(fornecedor)

        # Preço: reconstrói o corrompido a partir de custo x markup
        custo = num(g(r, "Preço de custo")) or 0
        markup = num(g(r, "Lucro"))
        venda_gravada = num(g(r, "Preço de venda")) or 0
        custo_ok = 0 < custo <= 10000
        if custo_ok and markup is not None:
            preco_venda = round(custo * (1 + markup / 100), 2)
        elif 0 < venda_gravada <= 10000:
            preco_venda = venda_gravada
        else:
            preco_venda = 0
        custo_final = custo if custo_ok else 0

        produtos.append({
            "codigoExterno": codigo,
            "descricao": descricao,
            "codigoBarras": barras,
            "unidade": unidade_clean(g(r, "Unidade")),
            "ncm": ncm_clean(g(r, "NCM")),
            "cest": code_str(g(r, "Cest")),
            "origem": code_str(g(r, "Origem")),
            "csosn": code_str(g(r, "Cst")),
            "icmsAliquota": num(g(r, "Icms")),
            "fabricante": extrair_marca(descricao, g(r, "Fabricante"), marcas),
            "precoCusto": custo_final,
            "custoMedio": custo_final,
            "markup": markup or 0,
            "precoVenda": preco_venda,
            "saldoEstoque": num(g(r, "Quantidade")) or 0,
            "estoqueMinimo": num(g(r, "Estoque mínimo")) or 0,
            "fornecedor": fornecedor,
        })

    # Dedup por impressão digital (letras/números): mantém o de maior estoque
    melhor = {}
    for p in produtos:
        k = fingerprint(p["descricao"])
        if k not in melhor or (p["saldoEstoque"] or 0) > (melhor[k]["saldoEstoque"] or 0):
            melhor[k] = p
    dedup = list(melhor.values())
    com_marca = sum(1 for p in dedup if p["fabricante"])

    with open(saida, "w", encoding="utf-8") as f:
        json.dump({"produtos": dedup, "fornecedores": sorted(fornecedores)}, f, ensure_ascii=False)

    print(f"Produtos: {len(dedup)} (de {len(data)} linhas) | Fornecedores: {len(fornecedores)}")
    print(f"Com marca preenchida: {com_marca} ({100*com_marca/len(dedup):.0f}%) | "
          f"dup. de código: {pulados_dup} | dup. de nome removidas: {len(produtos)-len(dedup)}")


if __name__ == "__main__":
    main()
